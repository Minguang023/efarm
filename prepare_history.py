"""Read the supplied Excel workbook into an auditable, portable import bundle.
Usage: python prepare_history.py workbook.xlsx history/records.json
The original workbook is never modified.
"""
import hashlib
import json
import re
import sys
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from uuid import uuid5, NAMESPACE_URL
import openpyxl

MONTHS = dict(JAN=1,FEB=2,MAR=3,APR=4,MAY=5,JUNE=6,JULY=7,AUG=8,SEPT=9,OCT=10,NOV=11,DEC=12)

def parsed_date(value):
    if isinstance(value, datetime): return value.date().isoformat()
    if isinstance(value, date): return value.isoformat()
    value = str(value).strip()
    for fmt in ('%d-%m-%Y','%d/%m/%Y','%d-%m-%y','%d/%m/%y','%Y-%m-%d'):
        try: return datetime.strptime(value,fmt).date().isoformat()
        except ValueError: pass
    return None

def sen(value):
    if not isinstance(value,(int,float,Decimal)): return None
    return int((Decimal(str(value))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))

def code_text(cell):
    value=cell.value
    if value is None: return ''
    if isinstance(value,(int,float)) and float(value).is_integer():
        raw=str(int(value))
        if re.fullmatch('0+',cell.number_format): raw=raw.zfill(len(cell.number_format))
        return raw
    return str(value).strip()

def category(description):
    # Only explicit category descriptions are classified. No colour/supplier guesses.
    text=description.strip().upper()
    if text in ('SALARY','ACCOUNTING SERVICES','EPF-SOCSC','EPF-SOCSO','EPF','SOCSO'):
        return 'Salary/EPF/Accounting'
    if text == 'RENTAL': return 'Rental'
    return None

def prepare(path):
    path=Path(path)
    w=openpyxl.load_workbook(path,data_only=True)
    result=dict(schema_version=2,source_name=path.name,source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                sales=[],expenses=[],summaries=[],blank_sales=[],issues=[])
    for sheet in w:
        match=re.search(r'(JAN|FEB|MAR|APR|MAY|JUNE|JULY|AUG|SEPT|OCT|NOV|DEC)\s+(20\d{2})',sheet.title)
        if not match: continue
        month=f'{match[2]}-{MONTHS[match[1]]:02d}-01'
        if 'DAILY' in sheet.title:
            for row in sheet.iter_rows(min_row=3):
                day=parsed_date(row[0].value)
                if not day: continue
                vals=[c.value for c in row[1:4]]
                source=dict(sheet=sheet.title,row=row[0].row,raw=vals,raw_date=str(row[0].value))
                if day[:7]!=month[:7]:
                    result['issues'].append(dict(sheet=sheet.title,row=row[0].row,code='',issue='Sales date normalized to the sheet year/month; original date retained'))
                    day=date(int(match[2]),MONTHS[match[1]],date.fromisoformat(day).day).isoformat()
                if all(v is None for v in vals):
                    result['blank_sales'].append(dict(sale_date=day,**source));continue
                amounts=[sen(v) or 0 for v in vals]
                if any(a<0 for a in amounts): raise ValueError(f'Negative sales at {sheet.title}:{row[0].row}')
                result['sales'].append(dict(sale_date=day,sales_1_sen=amounts[0],sales_2_sen=amounts[1],
                                           e_wallet_sen=amounts[2],source=source))
        elif 'COST' in sheet.title:
            for row in sheet.iter_rows(min_row=2):
                raw=[c.value for c in row[:5]]
                if not any(v is not None for v in raw): continue
                if any(str(raw[i]).strip().upper() in ('TOTAL','TOTAL COST','TOTAL EXPENSES') for i in (0,3)): continue
                if raw[0] is None and raw[1] is None and raw[3] is None and raw[4] is None:
                    result['issues'].append(dict(sheet=sheet.title,row=row[0].row,code='',issue='Non-record note retained: '+str(raw[2])))
                    continue
                if not raw[1] and not raw[2] and not raw[3]: continue
                issues=[]
                day=parsed_date(raw[0]);amount=sen(raw[4]);code=code_text(row[2])
                if not day: issues.append('Invalid or missing invoice date')
                elif day[:7]>month[:7]: issues.append('Invoice date is later than its reporting month')
                if amount is None: issues.append('Missing or nonnumeric amount; excluded from totals until corrected')
                if amount is not None and amount<0: issues.append('Negative source amount retained')
                if isinstance(raw[2],(int,float)) and abs(raw[2])>=1e15:
                    issues.append('Long numeric invoice code: Excel may already have lost digit precision')
                record_id=str(uuid5(NAMESPACE_URL,f'efarm/{sheet.title.strip()}/row/{row[0].row}'))
                desc=str(raw[3] or '').strip()
                result['expenses'].append(dict(id=record_id,invoice_date=day,reporting_month=month,
                    invoice_code=code,description=desc,category=category(desc),amount_sen=amount,
                    source_sheet=sheet.title,source_row=row[0].row,legacy_code=str(raw[1] or ''),
                    source_json=json.dumps(raw,default=str,ensure_ascii=False),review_note='; '.join(issues)))
                for issue in issues: result['issues'].append(dict(sheet=sheet.title,row=row[0].row,code=str(raw[1]),issue=issue))
        elif 'SUMMARY' in sheet.title:
            debit={};credit={}
            for row in sheet.iter_rows():
                label=str(row[0].value or '').strip()
                if not label or label.upper().startswith(('SUMMARY','PROFIT')): continue
                if len(row)>1 and sen(row[1].value) is not None: debit[label]=sen(row[1].value)
                if len(row)>2 and sen(row[2].value) is not None: credit[label]=sen(row[2].value)
            result['summaries'].append(dict(reporting_month=month,source_sheet=sheet.title,debit=debit,credit=credit))
    for summary in result['summaries']:
        month=summary['reporting_month'][:7]
        summary['detail_sales_sen']=sum(sum(r[k] for k in ('sales_1_sen','sales_2_sen','e_wallet_sen')) for r in result['sales'] if r['sale_date'][:7]==month)
        summary['detail_expenses_sen']=sum(r['amount_sen'] or 0 for r in result['expenses'] if r['reporting_month'][:7]==month)
    return result

if __name__=='__main__':
    bundle=prepare(sys.argv[1]);Path(sys.argv[2]).write_text(json.dumps(bundle,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(json.dumps({k:len(bundle[k]) for k in ('sales','expenses','summaries','blank_sales','issues')}))
